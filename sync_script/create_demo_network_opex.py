#!/usr/bin/env python3
"""Заполнить бюджет OPEX по контракту в реестре сетей demo-БД.

Бюджет OPEX — второй механизм инвестиций реестра, рядом с процентом от
товарооборота: согласованная сумма за услуги сети по пяти статьям договора
(см. миграцию 031 и services/network_opex_service.go). Демо-реестр его не
заполняет — вкладка «Инвестиции OPEX» и столбец OPEX витрины стоят пустыми.

Правило наполнения:

- Годовой бюджет сети — доля от 1 до 5 % годового товарооборота. Товарооборот —
  сумма плана брендов за год (`tbl_NetworkPlans`, `brand_as IS NOT NULL`); строка
  валового пула (`brand_as IS NULL`) — общий объём контракта, а не бренд, и в
  базу не входит. План, а не факт: бюджет по договору согласуют до отгрузок, и
  у незакрытого года факта на весь год нет.
- Процент — свойство сети и на все годы один: бюджет услуг в контракте меняют
  редко, а товарооборот года уже даёт разницу между годами.
- Бюджет делится между брендами плана по их доле в товарообороте — таблица
  требует бренд, и бренд обязан быть в плане года (иначе `SaveNetworkOpexBudgets`
  такую строку не примет).
- Внутри бренда бюджет делится по статьям с весами сети; часть статей у сети
  отсутствует — не каждая сеть ведёт карточки товара на сайте.
- По году равномерно: квартал — ровно четверть, округлённая до рубля, как её
  согласовал бы КАМ. Квартал раскладывается на месяцы тем же правилом, что и
  сервер (services.NetworkOpexMonthlyRows): копейки, остаток в последний месяц.
- Обе базы НДС, как и на сервере: с НДС — введённая сумма, без НДС — по ставке
  квартала из `tbl_NetworkPeriods` (незаведённый квартал берёт карточку сети).

Факта у бюджета нет по модели данных (README, «Бюджет OPEX по контракту»):
единственное число — согласованный бюджет, и «факт равен плану» за прошлые
периоды выполняется по построению.

Все случайности устойчивы (StableSynthetic): повторный запуск даёт те же
числа. Трогаются только годы из `--years`; бюджет других лет остаётся.
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import sys
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Sequence

from dotenv import dotenv_values

from create_demo_ecom_sales import pick_decimal, unit_fraction
from create_demo_promo_db import (
    StableSynthetic,
    clean_text,
    connect,
    execute_many,
    fetch_dicts,
    quantize_money,
)


RESET_CONFIRMATION = "RESET_DEMO_NETWORK_OPEX"
UPDATED_BY = "demo-opex"

# Статьи в порядке показа — копия services.opexArticles и CK миграции 031.
ARTICLES = (
    ("ntz_bdn", "Поддержание НТЗ/БДН"),
    ("display", "Выкладка"),
    ("reports", "Предоставление отчётов"),
    ("fixed_promo", "Акция с фиксированной стоимостью"),
    ("product_card", "Размещение карточки товара на сайте сети"),
)
ARTICLE_LABELS = dict(ARTICLES)

# Доля бюджета OPEX от годового товарооборота, в процентах.
OPEX_PCT_BOUNDS = ("1.00", "5.00")
# Вероятность, что статьи у сети нет вовсе; статей не меньше MIN_ARTICLES.
ARTICLE_ABSENT_SHARE = 0.20
MIN_ARTICLES = 2

ONE_RUB = Decimal("1")


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Заполнить бюджет OPEX по контракту в реестре сетей demo-БД."
    )
    parser.add_argument("--env-file", type=Path, default=Path(".env"))
    parser.add_argument("--target-server", default="127.0.0.1,1434")
    parser.add_argument("--target-db", default=None)
    parser.add_argument("--years", type=int, nargs="+", default=[2025, 2026])
    parser.add_argument("--dry-run", action="store_true", help="Собрать строки и таблицу, в базу не писать")
    parser.add_argument("--export", type=Path, default=None, help="Куда выгрузить таблицу (.xlsx или .csv)")
    parser.add_argument("--replace", action="store_true")
    parser.add_argument("--confirm", default="")
    return parser.parse_args(argv)


def require_demo_target(target_db: str) -> None:
    if "demo" not in target_db.casefold():
        raise ValueError("Имя целевой базы обязано содержать 'demo'")


# ─── Чтение реестра ────────────────────────────────────────────────────────

def fetch_networks(cursor) -> list[dict[str, Any]]:
    return fetch_dicts(
        cursor,
        """
        SELECT id, name, kam, is_active, vat_included, vat_rate
        FROM dbo.tbl_Networks ORDER BY id
        """,
    )


def fetch_period_vat(cursor, years: Sequence[int]) -> dict[tuple[int, int, int], tuple[bool, Decimal]]:
    """НДС квартала: (network_id, year, quarter) → (vat_included, vat_rate)."""
    rows = fetch_dicts(
        cursor,
        f"""
        SELECT network_id, [year], [quarter], vat_included, vat_rate
        FROM dbo.tbl_NetworkPeriods
        WHERE [year] IN ({",".join("?" for _ in years)})
        """,
        tuple(years),
    )
    return {
        (int(row["network_id"]), int(row["year"]), int(row["quarter"])):
            (bool(row["vat_included"]), Decimal(str(row["vat_rate"] or 0)))
        for row in rows
    }


def fetch_brand_turnover(cursor, years: Sequence[int]) -> dict[tuple[int, int], dict[str, Decimal]]:
    """Годовой план бренда: (network_id, year) → {brand_as: plan_rub}."""
    rows = fetch_dicts(
        cursor,
        f"""
        SELECT network_id, [year], LTRIM(RTRIM(brand_as)) AS brand_as, SUM(plan_rub) AS plan_rub
        FROM dbo.tbl_NetworkPlans
        WHERE [year] IN ({",".join("?" for _ in years)})
          AND brand_as IS NOT NULL AND LTRIM(RTRIM(brand_as)) <> ''
        GROUP BY network_id, [year], LTRIM(RTRIM(brand_as))
        """,
        tuple(years),
    )
    result: dict[tuple[int, int], dict[str, Decimal]] = defaultdict(dict)
    for row in rows:
        brand = clean_text(row.get("brand_as"))
        plan = Decimal(str(row.get("plan_rub") or 0))
        if brand and plan > 0:
            result[(int(row["network_id"]), int(row["year"]))][brand] = plan
    return result


# ─── Правило ───────────────────────────────────────────────────────────────

def network_opex_pct(synthetic: StableSynthetic, name: str) -> Decimal:
    """Доля бюджета от товарооборота — свойство сети, одна на все годы."""
    return pick_decimal(synthetic, "opex-pct", name, *OPEX_PCT_BOUNDS).quantize(
        Decimal("0.01"), rounding=ROUND_HALF_UP
    )


def article_weights(synthetic: StableSynthetic, name: str) -> dict[str, Decimal]:
    """Веса статей сети, сумма — единица; отсутствующих статей в словаре нет."""
    raw = {
        code: Decimal(str(round(unit_fraction(synthetic, "opex-article-weight", f"{name}|{code}"), 6)))
        + Decimal("0.05")
        for code, _ in ARTICLES
    }
    present = [
        code for code, _ in ARTICLES
        if unit_fraction(synthetic, "opex-article-present", f"{name}|{code}") >= ARTICLE_ABSENT_SHARE
    ]
    if len(present) < MIN_ARTICLES:
        present = sorted(raw, key=lambda code: raw[code], reverse=True)[:MIN_ARTICLES]
    total = sum(raw[code] for code in present)
    return {code: raw[code] / total for code in present}


def split_kopecks(amount: Decimal) -> tuple[Decimal, Decimal, Decimal]:
    """Квартал на три месяца до копейки, остаток — в последний: как splitOpexKopecks."""
    total = int(quantize_money(amount) * 100)
    # Целочисленное деление Go усекает к нулю; у Python — к минус бесконечности,
    # поэтому знак снимается до деления, чтобы отрицательный бюджет лёг так же.
    sign = -1 if total < 0 else 1
    share = sign * (abs(total) // 3)
    return (
        Decimal(share) / 100,
        Decimal(share) / 100,
        Decimal(total - 2 * share) / 100,
    )


def net_rub(gross: Decimal, vat_included: bool, vat_rate: Decimal) -> Decimal:
    """База «без НДС» — как services.NetRub."""
    if not vat_included or vat_rate <= 0:
        return quantize_money(gross)
    return quantize_money(gross / (1 + vat_rate / 100))


def build_opex_rows(
    synthetic: StableSynthetic,
    networks: list[dict[str, Any]],
    turnover: dict[tuple[int, int], dict[str, Decimal]],
    period_vat: dict[tuple[int, int, int], tuple[bool, Decimal]],
    years: Sequence[int],
) -> tuple[list[tuple[Any, ...]], list[dict[str, Any]]]:
    """Месячные строки для tbl_NetworkOpexBudgets и квартальные ячейки для таблицы.

    Строки — в порядке сети, года, бренда, статьи, месяца. Ячейка с нулевым
    кварталом не заводится: ноль в базе означает «согласован нулевой бюджет»,
    а здесь это просто слишком маленький бренд.
    """
    rows: list[tuple[Any, ...]] = []
    cells: list[dict[str, Any]] = []
    for network in networks:
        network_id = int(network["id"])
        name = clean_text(network.get("name"))
        pct = network_opex_pct(synthetic, name)
        weights = article_weights(synthetic, name)
        card_vat = (bool(network["vat_included"]), Decimal(str(network["vat_rate"] or 0)))

        for year in years:
            brands = turnover.get((network_id, year), {})
            total_turnover = sum(brands.values(), Decimal(0))
            if total_turnover <= 0:
                continue
            annual_budget = total_turnover * pct / Decimal(100)

            for brand in sorted(brands):
                brand_budget = annual_budget * brands[brand] / total_turnover
                for code, _ in ARTICLES:
                    weight = weights.get(code)
                    if weight is None:
                        continue
                    quarter_amount = (brand_budget * weight / 4).quantize(ONE_RUB, rounding=ROUND_HALF_UP)
                    if quarter_amount == 0:
                        continue
                    for quarter in (1, 2, 3, 4):
                        vat_included, vat_rate = period_vat.get((network_id, year, quarter), card_vat)
                        gross_months = split_kopecks(quarter_amount)
                        net_months = split_kopecks(net_rub(quarter_amount, vat_included, vat_rate))
                        for index in range(3):
                            rows.append((
                                network_id, year, (quarter - 1) * 3 + 1 + index, brand, code,
                                gross_months[index], net_months[index], UPDATED_BY,
                            ))
                        cells.append({
                            "network_id": network_id, "network": name, "kam": clean_text(network.get("kam")),
                            "year": year, "quarter": quarter, "brand_as": brand, "article": code,
                            "amount_rub": quarter_amount,
                            "amount_rub_net": net_rub(quarter_amount, vat_included, vat_rate),
                            "turnover_rub": total_turnover, "pct": pct,
                        })
    return rows, cells


# ─── Таблица ───────────────────────────────────────────────────────────────

def summarize_networks(cells: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Сеть × год: товарооборот, процент, бюджет по статьям и кварталам."""
    buckets: dict[tuple[int, int], dict[str, Any]] = {}
    for cell in cells:
        key = (cell["network_id"], cell["year"])
        row = buckets.get(key)
        if row is None:
            row = {
                "network_id": cell["network_id"], "network": cell["network"], "kam": cell["kam"],
                "year": cell["year"], "turnover_rub": cell["turnover_rub"], "pct": cell["pct"],
                "budget_rub": Decimal(0), "budget_rub_net": Decimal(0),
                **{f"q{quarter}_rub": Decimal(0) for quarter in (1, 2, 3, 4)},
                **{f"{code}_rub": Decimal(0) for code, _ in ARTICLES},
                "brands": set(),
            }
            buckets[key] = row
        row["budget_rub"] += cell["amount_rub"]
        row["budget_rub_net"] += cell["amount_rub_net"]
        row[f"q{cell['quarter']}_rub"] += cell["amount_rub"]
        row[f"{cell['article']}_rub"] += cell["amount_rub"]
        row["brands"].add(cell["brand_as"])
    result = []
    for key in sorted(buckets):
        row = buckets[key]
        row["brands"] = len(row["brands"])
        row["effective_pct"] = (row["budget_rub"] / row["turnover_rub"] * 100).quantize(
            Decimal("0.01"), rounding=ROUND_HALF_UP
        )
        result.append(row)
    return result


SUMMARY_COLUMNS = (
    ("network_id", "ID сети"), ("network", "Сеть"), ("kam", "КАМ"), ("year", "Год"),
    ("turnover_rub", "Товарооборот (план брендов), руб"), ("pct", "Доля OPEX, %"),
    ("effective_pct", "Факт. доля после округления, %"), ("brands", "Брендов"),
    ("budget_rub", "Бюджет OPEX, руб с НДС"), ("budget_rub_net", "Бюджет OPEX, руб без НДС"),
    ("q1_rub", "Q1, руб"), ("q2_rub", "Q2, руб"), ("q3_rub", "Q3, руб"), ("q4_rub", "Q4, руб"),
    *((f"{code}_rub", f"{label}, руб") for code, label in ARTICLES),
)

CELL_COLUMNS = (
    ("network_id", "ID сети"), ("network", "Сеть"), ("year", "Год"), ("quarter", "Квартал"),
    ("brand_as", "Бренд"), ("article", "Код статьи"), ("article_label", "Статья"),
    ("amount_rub", "Сумма, руб с НДС"), ("amount_rub_net", "Сумма, руб без НДС"),
)


def export_table(path: Path, summary: list[dict[str, Any]], cells: list[dict[str, Any]]) -> None:
    detailed = [{**cell, "article_label": ARTICLE_LABELS[cell["article"]]} for cell in cells]
    if path.suffix.casefold() == ".csv":
        with path.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.writer(handle, delimiter=";")
            writer.writerow(label for _, label in SUMMARY_COLUMNS)
            for row in summary:
                writer.writerow(row[key] for key, _ in SUMMARY_COLUMNS)
        return

    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheets = (
        ("Сети по годам", SUMMARY_COLUMNS, summary),
        ("Ячейки квартал×бренд×статья", CELL_COLUMNS, detailed),
    )
    for index, (title, columns, data) in enumerate(sheets):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = title
        sheet.append([label for _, label in columns])
        for row in data:
            sheet.append([
                float(row[key]) if isinstance(row[key], Decimal) else row[key]
                for key, _ in columns
            ])
        sheet.freeze_panes = "A2"
        for column_index, (key, label) in enumerate(columns, start=1):
            letter = get_column_letter(column_index)
            sheet.column_dimensions[letter].width = max(12, min(40, len(label) + 2))
            if key.endswith("_rub"):
                for cell in sheet[letter][1:]:
                    cell.number_format = "#,##0.00"
    workbook.save(path)


# ─── Загрузка ──────────────────────────────────────────────────────────────

def count_opex_rows(cursor, years: Sequence[int]) -> int:
    cursor.execute(
        f"SELECT COUNT_BIG(*) FROM dbo.tbl_NetworkOpexBudgets WHERE [year] IN ({','.join('?' for _ in years)})",
        tuple(years),
    )
    return int(cursor.fetchone()[0])


def clear_opex_rows(cursor, years: Sequence[int]) -> None:
    cursor.execute(
        f"DELETE FROM dbo.tbl_NetworkOpexBudgets WHERE [year] IN ({','.join('?' for _ in years)})",
        tuple(years),
    )


def insert_opex_rows(cursor, rows: list[tuple[Any, ...]]) -> int:
    execute_many(
        cursor,
        """
        INSERT INTO dbo.tbl_NetworkOpexBudgets
            (network_id, [year], [month], brand_as, article, amount_rub, amount_rub_net, updated_by)
        VALUES (?,?,?,?,?,?,?,?)
        """,
        rows,
        batch_size=500,
    )
    return len(rows)


def verify_target(cursor, years: Sequence[int]) -> dict[str, Any]:
    placeholders = ",".join("?" for _ in years)

    # Бренд бюджета обязан быть в плане года — то же условие, что и у сервера.
    cursor.execute(
        f"""
        SELECT COUNT_BIG(*) FROM dbo.tbl_NetworkOpexBudgets o
        WHERE o.[year] IN ({placeholders}) AND NOT EXISTS (
            SELECT 1 FROM dbo.tbl_NetworkPlans p
            WHERE p.network_id = o.network_id AND p.[year] = o.[year]
              AND LTRIM(RTRIM(p.brand_as)) = o.brand_as
        )
        """,
        tuple(years),
    )
    orphans = int(cursor.fetchone()[0])
    if orphans:
        raise RuntimeError(f"{orphans} строк бюджета OPEX ссылаются на бренд вне плана года")

    # Каждая ячейка — ровно три месяца своего квартала, а кварталы года равны.
    cursor.execute(
        f"""
        SELECT COUNT_BIG(*) FROM (
            SELECT network_id, [year], brand_as, article, (([month]-1)/3)+1 AS q, COUNT(*) AS months
            FROM dbo.tbl_NetworkOpexBudgets WHERE [year] IN ({placeholders})
            GROUP BY network_id, [year], brand_as, article, (([month]-1)/3)+1
        ) c WHERE months <> 3
        """,
        tuple(years),
    )
    broken = int(cursor.fetchone()[0])
    if broken:
        raise RuntimeError(f"{broken} ячеек бюджета OPEX без полных трёх месяцев")

    cursor.execute(
        f"""
        SELECT COUNT_BIG(*) FROM (
            SELECT network_id, [year], brand_as, article,
                   MIN(q_rub) AS min_q, MAX(q_rub) AS max_q, COUNT(*) AS quarters
            FROM (
                SELECT network_id, [year], brand_as, article, (([month]-1)/3)+1 AS q, SUM(amount_rub) AS q_rub
                FROM dbo.tbl_NetworkOpexBudgets WHERE [year] IN ({placeholders})
                GROUP BY network_id, [year], brand_as, article, (([month]-1)/3)+1
            ) q GROUP BY network_id, [year], brand_as, article
        ) c WHERE min_q <> max_q OR quarters <> 4
        """,
        tuple(years),
    )
    uneven = int(cursor.fetchone()[0])
    if uneven:
        raise RuntimeError(f"{uneven} ячеек бюджета OPEX распределены по кварталам неравномерно")

    # Доля от товарооборота — в границах правила, с запасом на округление до рубля.
    stats = fetch_dicts(
        cursor,
        f"""
        SELECT o.[year], COUNT(DISTINCT o.network_id) AS networks, COUNT_BIG(*) AS rows_,
               SUM(o.amount_rub) AS budget_rub, SUM(o.amount_rub_net) AS budget_rub_net,
               MIN(o.share_pct) AS min_pct, MAX(o.share_pct) AS max_pct
        FROM (
            SELECT b.[year], b.network_id, b.amount_rub, b.amount_rub_net,
                   SUM(b.amount_rub) OVER (PARTITION BY b.network_id, b.[year]) * 100.0 / t.plan_rub AS share_pct
            FROM dbo.tbl_NetworkOpexBudgets b
            JOIN (
                SELECT network_id, [year], SUM(plan_rub) AS plan_rub
                FROM dbo.tbl_NetworkPlans WHERE brand_as IS NOT NULL
                GROUP BY network_id, [year]
            ) t ON t.network_id = b.network_id AND t.[year] = b.[year]
            WHERE b.[year] IN ({placeholders})
        ) o GROUP BY o.[year] ORDER BY o.[year]
        """,
        tuple(years),
    )
    low, high = (Decimal(bound) for bound in OPEX_PCT_BOUNDS)
    for item in stats:
        if Decimal(str(item["min_pct"])) < low - Decimal("0.05") or Decimal(str(item["max_pct"])) > high + Decimal("0.05"):
            raise RuntimeError(
                f"Доля OPEX {item['year']} года вне границ {low}–{high} %: "
                f"{item['min_pct']}…{item['max_pct']}"
            )
    return {
        str(item["year"]): {
            "networks": int(item["networks"]),
            "rows": int(item["rows_"]),
            "budget_rub": str(quantize_money(Decimal(str(item["budget_rub"])))),
            "budget_rub_net": str(quantize_money(Decimal(str(item["budget_rub_net"])))),
            "share_pct": f"{Decimal(str(item['min_pct'])):.2f}…{Decimal(str(item['max_pct'])):.2f}",
        }
        for item in stats
    }


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    env = dotenv_values(args.env_file)
    password = env.get("SA_PASSWORD") or os.getenv("SA_PASSWORD")
    target_db = args.target_db or env.get("DEMO_DB_NAME") or "local_project_demo_db"
    salt = env.get("DEMO_MAPPING_SALT") or env.get("JWT_SECRET") or os.getenv("DEMO_MAPPING_SALT")
    if not password:
        raise RuntimeError("SA_PASSWORD не задан")
    if not salt:
        raise RuntimeError("DEMO_MAPPING_SALT или JWT_SECRET не задан")
    if args.replace and args.confirm != RESET_CONFIRMATION:
        raise ValueError(f"Для --replace требуется --confirm {RESET_CONFIRMATION}")
    require_demo_target(target_db)
    years = sorted(set(args.years))

    synthetic = StableSynthetic(salt)
    target = connect(args.target_server, target_db, password, readonly=args.dry_run)
    cursor = target.cursor()

    try:
        networks = fetch_networks(cursor)
        turnover = fetch_brand_turnover(cursor, years)
        if not networks or not turnover:
            raise RuntimeError("В demo-БД нет реестра сетей; сначала выполните make demo-registry-load")
        period_vat = fetch_period_vat(cursor, years)

        rows, cells = build_opex_rows(synthetic, networks, turnover, period_vat, years)
        summary = summarize_networks(cells)
        if args.export:
            export_table(args.export, summary, cells)

        existing = count_opex_rows(cursor, years)
        report: dict[str, Any] = {
            "status": "dry-run" if args.dry_run else "ok",
            "target_database": target_db,
            "years": years,
            "existing_rows_in_years": existing,
            "prepared": {
                "networks": len({cell["network_id"] for cell in cells}),
                "network_years": len(summary),
                "cells": len(cells),
                "monthly_rows": len(rows),
                "budget_rub": str(sum((row["budget_rub"] for row in summary), Decimal(0))),
            },
            "export": str(args.export) if args.export else None,
        }
        if args.dry_run:
            print(json.dumps(report, ensure_ascii=False, indent=2))
            return 0

        if existing and not args.replace:
            raise RuntimeError(
                f"Бюджет OPEX за {years} уже заведён ({existing} строк). "
                "Используйте --replace с защитной фразой."
            )
        if existing:
            clear_opex_rows(cursor, years)
        inserted = insert_opex_rows(cursor, rows)
        try:
            verification = verify_target(cursor, years)
        except Exception:
            target.rollback()
            raise
        target.commit()

        report["inserted"] = inserted
        report["verification"] = verification
        print(json.dumps(report, ensure_ascii=False, indent=2))
        return 0
    except Exception:
        target.rollback()
        raise
    finally:
        cursor.close()
        target.close()


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as error:
        print(f"Ошибка наполнения бюджета OPEX demo-БД: {error}", file=sys.stderr)
        raise SystemExit(1)
